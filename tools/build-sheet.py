"""人間向けスプレッドシート(xlsx)を作る。決まりは knowledge/sheet-style.md。
各タブは GitHub の CSV を IMPORTDATA で読むだけ。列幅は今の CSV の中身から自動で決める。
使い方: python3 tools/build-sheet.py knowhow out.xlsx  /  python3 tools/build-sheet.py learning out.xlsx"""
import csv,math,sys,unicodedata,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
RAW="https://raw.githubusercontent.com/mukkii-game/Perfect_Dev_Environment/main/knowledge/tables/"
LOCAL="knowledge/tables/"
MAXPX,PAD,MINPX=500,14,36
def px(s):
  return sum(14 if unicodedata.east_asian_width(ch) in "WFA" else 7.5 for ch in str(s))+PAD
def width(vals):
  w=sorted(px(v) for v in vals); M=w[-1]
  if M<=MAXPX:
    rest=w[-4] if len(w)>4 else 0          # 数個(3つ)だけ長いなら 2 行で済む幅まで狭める
    t=max(rest,math.ceil((M-PAD)/2)+PAD) if len(w)>4 and w[-4]<M*0.6 else M
  else:
    t=min(MAXPX,math.ceil((M-PAD)/2)+PAD)   # 2 行に収める(無理なら上限)
  return max(MINPX,t)/7                     # xlsx の幅単位(≒7px)
RANK={"S":"C6EFCE","A":"DDEBF7","B":"FFF2CC","C":"F2F2F2","外した":"D9D9D9"}
BOOKS={
 "knowhow":("説明",["開発ノウハウ一覧(閲覧用)","原本は GitHub の knowledge/tables/*.csv。各タブは自動で読み込んで表示するだけ(初回だけ各タブの A1 で「アクセスを許可」)",
   "道具: 分野→カテゴリ→工程の順(BGMが欲しい→拾うか作るか、の順で引ける)。区分: 優先/条件付き/外した。AI操作: ◎AIだけで完結 ○MCP・APIでAIが操作 △人間の手が要る。ランク: S 迷ったらこれ/A 第一候補/B 補助/C ほぼ使わない",
   "一番右の「コメント(自由記入)」は自由に書いてよい。毎晩会議室が読んで取り込む(行の並びが変わるとずれるので、取り込まれたら消してよい)",
   "直したい時: 会議室に一言、または GitHub で CSV を編集。表の中身を直接書き換えても次の更新で消えます",
   "どこに何があるか: 「置き場所」タブ"],
  [("置き場所","storage.csv","A",{"GitHub":"DDEBF7","Google":"C6EFCE","Claude":"FFF2CC","手元":"F2F2F2"}),
   ("心構え","mindset.csv","B","rank"),("AI","ai.csv","C","rank"),("道具","assets.csv","D","rank"),
   ("ジャンル","genres.csv","J",{"◎":"C6EFCE","○":"DDEBF7","△":"FFF2CC","×":"D9D9D9"}),
   ("手続き","procedures.csv","H",{"実績あり":"C6EFCE","要確認":"FFF2CC"}),
   ("使った素材","used-assets.csv","A",{"★":"C6EFCE","✕":"D9D9D9"}),
   ("自動化","automations.csv","B",{"全自動":"C6EFCE","半自動":"DDEBF7","—":"F2F2F2"}),
   ("AIへの指示","rules.csv","C",{}),
   ("TIPS","tips.csv","B",{}),
   ("用語集","glossary.csv","B",{})]),
 "learning":("はじめに",["Mukkii 育成プログラム","目標: AI駆動で毎日ゲームを作り、プラットフォーム化する世界的プロ(Lv1〜7。定義は topics/10)",
   "進め方: 毎週月曜の週次レビューで「今週の課題」が1つ出る → 作品で実践 → 会議室に「テストして <単元ID>」→ 合格で日付が入る",
   "タブ: カリキュラム(単元一覧) / 進捗(毎回の記録) / テスト(出題と結果) / 修了チェック(Lvごとの昇級条件) / 関連リンク",
   "一番右の「コメント(自由記入)」に感想・質問・要望を書いてよい。毎晩会議室が読んで取り込み、カリキュラムを調整する"],
  [("カリキュラム","learning/curriculum.csv","I",{"合格":"C6EFCE","挑戦中":"FFF2CC"}),
   ("進捗","learning/progress.csv","A",{}),("テスト","learning/tests.csv","F",{"合格":"C6EFCE","不合格":"FCE4D6"}),
   ("修了チェック","learning/checklist.csv","E",{"昇級":"C6EFCE"}),("関連リンク","learning/links.csv","A",{})])}
def build(key,out):
  title,intro,tabs=BOOKS[key]
  wb=openpyxl.Workbook();ws=wb.active;ws.title=title
  for line in intro: ws.append([line])
  ws["A1"].font=Font(bold=True,size=14);ws.column_dimensions["A"].width=width(intro[1:])
  al=Alignment(vertical="center",wrap_text=True)
  ws.column_dimensions["A"].alignment=al
  for name,f,rc,mode in tabs:
    rows=list(csv.reader(open(LOCAL+f,encoding="utf-8")));n=len(rows[0]);last=len(rows)+20
    s=wb.create_sheet(name);s["A1"]=f'=IMPORTDATA("{RAW}{f}")';s.freeze_panes="A2"
    for i in range(n):
      d=s.column_dimensions[L(i+1)];d.width=width([r[i] for r in rows if i<len(r)]);d.alignment=al   # 列ごとの書式(全行に効く)
    cm=L(n+1);s[f"{cm}1"]="コメント(自由記入)";d=s.column_dimensions[cm];d.width=240/7;d.alignment=al
    d.fill=PatternFill("solid",fgColor="FFFDE7")
    for c in range(1,n+2): s.cell(1,c).font=Font(bold=True);s.cell(1,c).alignment=al
    if rows[0][-1]=="わかった":   # 理解チェック: 右端に ☑ 欄。毎晩会議室が CSV の「わかった」へ写して消す
      ck=L(n+2);s[f"{ck}1"]="わかった(☑)";s.column_dimensions[ck].width=12
    s.auto_filter.ref=f"A1:{L(n+1)}{last}"   # 1 行目にフィルタ(並べ替え・絞り込み)
    R=f"A2:{L(n)}{last}"
    if mode=="rank":
      for k,v in RANK.items(): s.conditional_formatting.add(R,FormulaRule(formula=[f'${rc}2="{k}"'],fill=PatternFill("solid",fgColor=v)))
    else:
      for k,v in mode.items(): s.conditional_formatting.add(R,FormulaRule(formula=[f'LEFT(${rc}2,{len(k)})="{k}"'],fill=PatternFill("solid",fgColor=v)))
  wb.save(out)
if __name__=="__main__": build(sys.argv[1],sys.argv[2])
